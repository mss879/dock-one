#!/usr/bin/env python3
"""
Builds the real-inventory migrations from the client's handoff workbook
(Dock_One_Combined_Inventory.xlsx: sheets "Combined Inventory" and "Brand Content").

  supabase/migrations/34_seed_inventory.sql        products, variants, prices, stock, sourcing notes
  supabase/migrations/35_seed_inventory_costs.sql  cost + dealer prices (git-ignored: the repo is public)

Every number (selling price, cost, dealer price, quantity) is read from the workbook cell; nothing
is typed in here. What this file DOES decide is the curation the workbook leaves to us:
  · which rows are one product with variants (rows that differ only by colour / size / capacity /
    switch / GPU, "consecutive within each product family" per the workbook),
  · the storefront department each sheet category goes into (DEPARTMENTS),
  · a tidy product name + one-line spec subtitle restating the row's own text (the full original
    text is kept per variant in variant_sourcing.source_description),
  · attributes.specs — only facts written in the row (or the matched brand title),
  · attributes.highlights — only the brand pages' short verbatim excerpts ("Short verbatim excerpt").
Price ranges ("215,000-220,000") become the LOWER figure; the range is kept in variant_sourcing.

Run:  python3 scripts/db/inventory/build_seed.py ~/Downloads/Dock_One_Combined_Inventory.xlsx
"""

import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
OUT_SEED = ROOT / "supabase/migrations/34_seed_inventory.sql"
OUT_COSTS = ROOT / "supabase/migrations/35_seed_inventory_costs.sql"

# ── storefront departments (approved: group the sheet's 20 categories) ────────────────────────
# id: (name, tagline, scene, sort_order) — only NEW ids are inserted; the four existing ones keep
# the owner's copy (mice is renamed only if it still has the seeded name).
NEW_CATEGORIES = {
    "monitors": ("Monitors", "Gaming monitors, 25 to 32 inch", "night", 50),
    "audio": ("Audio", "Over-ear headphones", "violet", 60),
    "power-charging": ("Power & Charging", "Chargers, cables, wireless chargers & power banks", "lime", 70),
    "cameras": ("Cameras & Instax", "Instant cameras, photo printers & film", "paper", 80),
    "components": ("Networking & Components", "Wi-Fi adapters & desktop RAM", "paper", 90),
}
DEPARTMENTS = {
    "Laptops": "laptops", "Monitors": "monitors", "Keyboards": "keyboards", "Gaming Combos": "keyboards",
    "Mice": "mice", "Mousepads": "mice", "Instant Cameras": "cameras", "Photo Printers": "cameras",
    "Instant Film": "cameras", "Chargers": "power-charging", "Cables": "power-charging",
    "Wireless Chargers": "power-charging", "Power Banks": "power-charging", "Headphones": "audio",
    "Internal SSD": "storage", "External HDD": "storage", "External SSD": "storage",
    "USB Flash Drives": "storage", "Network Adapter": "components", "Internal Storage": "components",
}
# facet tag per sheet category (the RAM row is filed as "Internal Storage" in the sheet; it is RAM)
TAGS = {"Internal Storage": "ram", "Network Adapter": "wi-fi adapter"}

DEMO_SLUGS = [
    "vanta-g15-gaming-laptop", "aeroslim-14-ultrabook", "forge-studio-16-creator-laptop",
    "campus-13-everyday-laptop", "bolt-x-portable-ssd-1tb", "atlas-slim-external-hdd-2tb",
    "duolink-flash-drive-128gb", "vault-desktop-backup-drive-8tb",
    "kairo-75-wireless-mechanical-keyboard", "onyx-pro-full-size-rgb-keyboard",
    "feather-slim-wireless-keyboard", "volt-60-compact-keyboard-acid-lime",
    "glide-mx-ergonomic-wireless-mouse", "aero-lite-gaming-mouse-58g",
    "grip-vertical-ergonomic-mouse", "pebble-go-travel-mouse",
]


def P(slug, brand, name, rows, *, subtitle=None, specs=None, use_cases=None, content=None, notes=None):
    """A product. rows = [(sheet_row, variant_name, option_values)] in selector order."""
    return dict(slug=slug, brand=brand, name=name, rows=rows, subtitle=subtitle, specs=specs or {},
                use_cases=use_cases or [], content=content, notes=notes)


def one(row, name="Standard", **options):
    return [(row, name, options)]


GB = lambda n: n  # noqa: E731 — readability: storage in decimal GB (1TB = 1000)

# content = the "Source item / match" key of the Brand Content row this product matches.
PRODUCTS = [
    # ── Laptops ────────────────────────────────────────────────────────────────────────────────
    P("hp-laptop-15-fd0133wm", "HP", "HP Laptop 15-fd0133wm", one(6),
      subtitle='Core i3-N305 · 8GB · 256GB SSD · 15.6" FHD · Windows 11',
      specs=dict(cpu="Intel Core i3-N305", ram_gb=8, storage_gb=256, storage_type="SSD", display="15.6-inch FHD", os="Windows 11"),
      content="15-fd0133wm"),
    P("hp-omnibook-16-by0015dx", "HP", "HP OmniBook 16-by0015dx", one(7),
      subtitle='Ryzen 5 · 8GB · 256GB SSD · 16" touch · Windows 11',
      specs=dict(cpu="AMD Ryzen 5", ram_gb=8, storage_gb=256, storage_type="SSD", display="16-inch touch", os="Windows 11"),
      content="16-by0015dx",
      notes='The stock list reads "Ryzen 5 40" — the processor model number looks cut off; confirm it before publishing.'),
    P("lenovo-v15-g4-83a1011jak", "Lenovo", "Lenovo V15 G4 (83A1011JAK)", one(8),
      subtitle='Core i3-1315U · 8GB · 256GB SSD · 15.6" FHD · DOS',
      specs=dict(cpu="Intel Core i3-1315U", ram_gb=8, storage_gb=256, storage_type="SSD", display="15.6-inch FHD", os="DOS"),
      content="83A1011JAK"),
    P("acer-aspire-go-15-ag15-42p-r917", "Acer", "Acer Aspire Go 15 AG15-42P-R917", one(9), content="AG15-42P"),
    P("dell-inspiron-ldc15255-a117blk-pus", "Dell", "Dell Inspiron LDC15255-A117BLK-PUS", one(10), content="LDC15255"),
    P("asus-vivobook-a1504va-bq541", "ASUS", "ASUS VivoBook A1504VA-BQ541", one(11),
      subtitle='Core 5 120U · 8GB DDR5 · 512GB SSD · 15.6" FHD',
      specs=dict(cpu="Intel Core 5 120U", ram_gb=8, storage_gb=512, storage_type="SSD", display="15.6-inch FHD"),
      content="A1504VA"),
    P("acer-aspire-lite-al15-53p-58sy", "Acer", "Acer Aspire Lite AL15-53P-58SY", one(12),
      subtitle='Core 5 120U · 16GB · 512GB SSD · 15.6"',
      specs=dict(cpu="Intel Core 5 120U", ram_gb=16, storage_gb=512, storage_type="SSD", display="15.6-inch"),
      content="AL15-53P"),
    P("asus-vivobook-x1607qa-ls56", "ASUS", "ASUS VivoBook X1607QA-LS56", one(13),
      subtitle='Snapdragon X X1-26-100 · 16GB DDR5 · 1TB SSD · 16"',
      specs=dict(cpu="Qualcomm Snapdragon X X1-26-100", ram_gb=16, storage_gb=1000, storage_type="SSD", display="16-inch"),
      content="X1607QA"),
    P("asus-vivobook-14-flip-tp3407sa", "ASUS", "ASUS VivoBook 14 Flip TP3407SA", one(14),
      subtitle='Core Ultra 7 256V · 16GB · 1TB SSD · 14" touch',
      specs=dict(cpu="Intel Core Ultra 7 256V", ram_gb=16, storage_gb=1000, storage_type="SSD", display="14-inch touch"),
      content="TP3407SA"),
    P("lenovo-loq-15ahp9", "Lenovo", "Lenovo LOQ 15AHP9", one(15),
      subtitle='Ryzen 7 7445HS · RTX 3050 6GB · 16GB DDR5 · 512GB SSD · 15.6"',
      specs=dict(cpu="AMD Ryzen 7 7445HS", gpu="NVIDIA GeForce RTX 3050 (6GB)", ram_gb=16, storage_gb=512, storage_type="SSD", display="15.6-inch"),
      content="15AHP9"),
    P("msi-cyborg-15-core-7-240h", "MSI", "MSI Cyborg 15 Gaming (Core 7 240H)",
      [(16, "RTX 5060 8GB", {"Graphics": "RTX 5060 8GB"}), (17, "RTX 5070 8GB", {"Graphics": "RTX 5070 8GB"})],
      subtitle='Core 7 240H · 16GB · 1TB SSD · 15.6"',
      specs=dict(cpu="Intel Core 7 240H", gpu="NVIDIA GeForce RTX 5060 (8GB)", ram_gb=16, storage_gb=1000, storage_type="SSD", display="15.6-inch"),
      use_cases=["gaming"], content="Core 7-240H"),
    P("msi-cyborg-15-i5-13420h", "MSI", "MSI Cyborg 15 (Core i5-13420H)", one(18),
      subtitle='Core i5-13420H · RTX 4050 6GB · 16GB · 512GB SSD · 15.6"',
      specs=dict(cpu="Intel Core i5-13420H", gpu="NVIDIA GeForce RTX 4050 (6GB)", ram_gb=16, storage_gb=512, storage_type="SSD", display="15.6-inch"),
      content="i5-13420H"),
    P("asus-tuf-a18-fa808uh-rs74", "ASUS", "ASUS TUF A18 FA808UH-RS74", one(19),
      subtitle='Ryzen 7 260 · RTX 5050 8GB · 16GB DDR5 · 1TB SSD · 18"',
      specs=dict(cpu="AMD Ryzen 7 260", gpu="NVIDIA GeForce RTX 5050 (8GB)", ram_gb=16, storage_gb=1000, storage_type="SSD", display="18-inch"),
      content="FA808UH"),
    # ── Monitors ───────────────────────────────────────────────────────────────────────────────
    P("kgaming-kg25dhi-280hz-25-gaming-monitor", "KGaming", 'KGaming KG25DHI 280Hz 25" Gaming Monitor', one(20, "25-inch", Size="25-inch"),
      subtitle="25-inch · 280Hz", specs=dict(display="25-inch", refresh_hz=280), use_cases=["gaming"], content="KG25DHI"),
    P("vision-forge-v1-4k-27-dual-mode-gaming-monitor", "KGaming", "Vision Forge V1 4K 27-inch Dual Mode Gaming Monitor", one(21, "27-inch", Size="27-inch"),
      subtitle="27-inch · 4K · Dual mode", specs=dict(display="27-inch 4K"), use_cases=["gaming"], content="Vision Forge V1 4K 27",
      notes="Brand not confirmed by the workbook (looked up on kgaming.co); confirm brand and model."),
    P("vision-forge-v1-4k-32-gaming-monitor", "KGaming", "Vision Forge V1 4K 32-inch Gaming Monitor", one(22, "32-inch", Size="32-inch"),
      subtitle="32-inch · 4K", specs=dict(display="32-inch 4K"), use_cases=["gaming"], content="Vision Forge V1 4K 32",
      notes="Brand not confirmed by the workbook (looked up on kgaming.co); confirm brand and model."),
    # ── Keyboards (incl. gaming combos) ────────────────────────────────────────────────────────
    P("kgaming-arctic-rift-tri-mode-gaming-keyboard", "KGaming", "KGaming Arctic Rift Tri-Mode Gaming Keyboard",
      one(23, "Cream Blue switch", Switch="Cream Blue"),
      subtitle="Tri-mode · Wired, Bluetooth & 2.4GHz",
      specs=dict(connectivity=["Wired", "Bluetooth", "2.4GHz wireless"]), use_cases=["gaming"], content="Arctic Rift"),
    P("iron-hawk-tri-mode-gaming-keyboard", "KGaming", "Iron Hawk Tri-Mode Gaming Keyboard",
      [(24, "Jelly Blue switch", {"Switch": "Jelly Blue"}), (25, "Jelly Pink switch", {"Switch": "Jelly Pink"})],
      subtitle="Tri-mode gaming keyboard", use_cases=["gaming"], content="Iron Hawk.*Blue",
      notes="HOLD: the stock list says Iron Hawk; the matching brand pages (C-MKB96-JB / -JP) say Arctic Rift. "
            "Keep the stock name until packaging or the supplier confirms the model."),
    P("4-in-1-gaming-combo-pack", "KGaming", "4-in-1 Gaming Combo Pack (Mouse, Mousepad, Keyboard, Headset)",
      [(26, "Blue switch", {"Switch": "Blue"}), (27, "Red switch", {"Switch": "Red"})],
      subtitle="Mouse · Mousepad · Keyboard · Headset", use_cases=["gaming"], content="4 in 1 Combo",
      notes="Brand not confirmed by the workbook (looked up on kgaming.co); confirm brand."),
    P("blackshadow-keyboard", "KGaming", "Blackshadow Keyboard",
      [(28, "Blue switch", {"Switch": "Blue"}), (29, "Red switch", {"Switch": "Red"})], content="Blackshadow",
      notes="Brand not confirmed by the workbook (looked up on kgaming.co); confirm brand."),
    P("logitech-mk270-wireless-keyboard-mouse-combo", "Logitech", "Logitech MK270 Wireless Keyboard and Mouse Combo", one(30),
      subtitle="Wireless · Full-size · 920-004519", specs=dict(layout="Full-size", connectivity=["Wireless"]), content="MK270"),
    P("logitech-mk295-silent-wireless-combo", "Logitech", "Logitech MK295 Silent Wireless Combo", one(31),
      subtitle="Wireless · Silent · 920-009801", specs=dict(connectivity=["Wireless"]), content="MK295"),
    # ── Mice & mousepads ───────────────────────────────────────────────────────────────────────
    P("logitech-m185-wireless-mouse", "Logitech", "Logitech M185 Wireless Mouse", one(32),
      subtitle="Wireless · 910-002235", specs=dict(connectivity=["Wireless"]), content="M185"),
    P("kgaming-aura-strike-rgb-wired-gaming-mouse", "KGaming", "KGaming Aura Strike RGB Wired Gaming Mouse", one(33),
      subtitle="RGB · Wired", specs=dict(connectivity=["Wired"], backlight="RGB"), use_cases=["gaming"], content="Aura Strike"),
    P("kgaming-space-star-rgb-wired-gaming-mouse", "KGaming", "KGaming Space Star RGB Wired Gaming Mouse", one(34),
      subtitle="RGB · Wired", specs=dict(connectivity=["Wired"], backlight="RGB"), use_cases=["gaming"], content="Space Star"),
    P("kgaming-glide-mousepad-cyborg", "KGaming", "KGaming Glide Mousepad Cyborg",
      [(35, "Large", {"Size": "Large"}), (36, "XL", {"Size": "XL"}), (37, "XXL", {"Size": "XXL"})],
      subtitle="Cyborg design · Large, XL or XXL", content="Glide Mousepad Cyborg Large"),
    P("kgaming-glide-mousepad-drakon", "KGaming", "KGaming Glide Mousepad Drakon",
      [(39, "Medium", {"Size": "Medium"}), (38, "Large", {"Size": "Large"})],
      subtitle="Drakon design · Medium or Large", content="Glide Mousepad Drakon Large",
      notes="Drakon Medium has no verified brand page (Drakon Large does)."),
    P("kgaming-glide-mousepad-white", "KGaming", "KGaming Glide Mousepad White", one(40, "Medium", Size="Medium"),
      subtitle="White · Medium", content="Glide Mousepad White Medium"),
    # ── Cameras & Instax ───────────────────────────────────────────────────────────────────────
    P("fujifilm-instax-mini-12", "Fujifilm", "Fujifilm instax mini 12",
      [(41, "Lilac Purple", {"Colour": "Lilac Purple"}), (42, "Mint Green", {"Colour": "Mint Green"}), (43, "Pastel Blue", {"Colour": "Pastel Blue"})],
      subtitle="Instant camera", content="MINI12"),
    P("fujifilm-instax-mini-se", "Fujifilm", "Fujifilm instax mini SE",
      [(44, "Blue", {"Colour": "Blue"}), (45, "Pink", {"Colour": "Pink"})], subtitle="Instant camera", content="MINI SE"),
    P("fujifilm-instax-mini-link-3", "Fujifilm", "Fujifilm instax mini Link 3",
      [(46, "Rose Pink", {"Colour": "Rose Pink"}), (47, "Sage Green", {"Colour": "Sage Green"})],
      subtitle="Smartphone photo printer", content="MINI LINK 3"),
    P("fujifilm-instax-mini-film", "Fujifilm", "Fujifilm instax mini Film",
      [(48, "10 sheets", {"Pack": "10 sheets"}), (49, "20 sheets", {"Pack": "20 sheets"})],
      subtitle="Instant film · 10 or 20 sheets", content="MINI FILM"),
    # ── Power & charging ───────────────────────────────────────────────────────────────────────
    P("anker-nano-charger-45w-a121d", "Anker", "Anker Nano Charger 45W (A121D)", one(50, "Blue", Colour="Blue"),
      subtitle="45W · USB-C", content="A121D"),
    P("anker-322-usb-c-to-usb-c-cable-3ft", "Anker", "Anker 322 USB-C to USB-C Cable (3ft)",
      [(51, "Black", {"Colour": "Black"}), (52, "White", {"Colour": "White"})], subtitle="USB-C to USB-C · 3ft", content="A81F5"),
    P("anker-nylon-usb-c-to-usb-c-cable-3ft-a8752", "Anker", "Anker Nylon USB-C to USB-C Cable (3ft, A8752)", one(53),
      subtitle="USB-C to USB-C · Nylon · 3ft", content="A8752"),
    P("anker-maggo-magnetic-wireless-charger", "Anker", "Anker MagGo Magnetic Wireless Charger",
      [(54, "White", {"Colour": "White"}), (55, "Black", {"Colour": "Black"})], subtitle="Magnetic wireless charger", content="B2568"),
    P("anker-533-power-bank-10k-30w", "Anker", "Anker 533 Power Bank (10K, 30W)", one(56),
      subtitle="10,000mAh · 30W", content="A1256"),
    P("anker-zolo-power-bank-25k-165w", "Anker", "Anker Zolo Power Bank (25K, 165W, Built-in Cable)", one(57),
      subtitle="25,000mAh · 165W · Built-in cable", content="A1695"),
    P("anker-737-power-bank-24k-140w", "Anker", "Anker 737 Power Bank (24K, 140W)", one(58),
      subtitle="24,000mAh · 140W", content="A1289"),
    # ── Audio ──────────────────────────────────────────────────────────────────────────────────
    P("anker-soundcore-life-q30", "Anker", "Anker Soundcore Life Q30", one(59, "Black", Colour="Black"),
      subtitle="Wireless headphones", content="LIFE Q30"),
    P("jbl-live-770nc", "JBL", "JBL Live 770NC", one(60, "Black", Colour="Black"),
      subtitle="Wireless over-ear · Adaptive noise cancelling", content="LIVE 770NC"),
    P("jbl-tune-730bt", "JBL", "JBL Tune 730BT", one(61, "Black", Colour="Black"),
      subtitle="Wireless over-ear", content="T730"),
    P("jbl-tune-770nc", "JBL", "JBL Tune 770NC", one(62, "Black", Colour="Black"),
      subtitle="Wireless over-ear · Adaptive noise cancelling", content="T770"),
    P("jbl-tune-780nc", "JBL", "JBL Tune 780NC", one(63, "Beige", Colour="Beige"),
      subtitle="Wireless over-ear · Noise cancelling", content="T780"),
    # ── Storage ────────────────────────────────────────────────────────────────────────────────
    P("lexar-ns100-sata-ssd", "Lexar", 'Lexar NS100 2.5" SATA SSD', one(64, "128GB", Capacity="128GB"),
      subtitle='Internal SSD · 2.5" SATA', specs=dict(capacity_gb=128, type="Internal SSD", interface='2.5" SATA'), content="NS100"),
    P("lexar-nm620-m2-nvme-ssd", "Lexar", "Lexar NM620 M.2 NVMe SSD", one(65, "256GB", Capacity="256GB"),
      subtitle="Internal SSD · M.2 NVMe Gen3 · 3500MB/s",
      specs=dict(capacity_gb=256, type="Internal SSD", interface="M.2 NVMe Gen3", speed_mbps=3500), content="NM620"),
    P("wd-my-passport-portable-hdd", "WD", "WD My Passport Portable HDD", one(66, "2TB", Capacity="2TB"),
      subtitle="External HDD", specs=dict(capacity_gb=2000, type="External HDD"), use_cases=["backup"], content="My Passport"),
    P("seagate-expansion-portable-hdd", "Seagate", "Seagate Expansion Portable HDD",
      [(67, "1TB", {"Capacity": "1TB"}), (68, "4TB", {"Capacity": "4TB"})],
      subtitle="External HDD · 1TB or 4TB", specs=dict(capacity_gb=1000, type="External HDD"), content="Seagate.*Expansion"),
    P("seagate-one-touch-portable-hdd", "Seagate", "Seagate One Touch Portable HDD",
      [(69, "1TB", {"Capacity": "1TB"}), (70, "2TB Black", {"Capacity": "2TB"})],
      subtitle="External HDD · 1TB or 2TB", specs=dict(capacity_gb=1000, type="External HDD"), content="Seagate.*One Touch"),
    P("sandisk-portable-ssd-1tb-e10", "SanDisk", "SanDisk Portable SSD (SDSSDE10)", one(71, "1TB", Capacity="1TB"),
      subtitle="Portable SSD · 600MB/s", specs=dict(capacity_gb=1000, type="Portable SSD", speed_mbps=600), content="SDSSDE10"),
    P("sandisk-portable-ssd-v3-e31", "SanDisk", "SanDisk Portable SSD V3 (E31)", one(72, "2TB", Capacity="2TB"),
      subtitle="Portable SSD · 1000MB/s", specs=dict(capacity_gb=2000, type="Portable SSD", speed_mbps=1000), content="E31"),
    P("sandisk-extreme-portable-ssd-ps5-e62p", "SanDisk", "SanDisk Extreme Portable SSD for PS5 (E62P)",
      [(73, "1TB", {"Capacity": "1TB"}), (74, "2TB", {"Capacity": "2TB"})],
      subtitle="Portable SSD · 1000MB/s · PS5 and PC", specs=dict(capacity_gb=1000, type="Portable SSD", speed_mbps=1000), content="E62P"),
    P("kingston-xs1000-portable-ssd", "Kingston", "Kingston XS1000 Portable SSD", one(75, "1TB", Capacity="1TB"),
      subtitle="Portable SSD", specs=dict(capacity_gb=1000, type="Portable SSD"), content="XS1000"),
    P("sandisk-ultra-dual-drive-go-usb-c", "SanDisk", "SanDisk Ultra Dual Drive Go USB Type-C",
      [(76, "32GB", {"Capacity": "32GB"}), (77, "64GB", {"Capacity": "64GB"}), (78, "128GB", {"Capacity": "128GB"}), (79, "256GB", {"Capacity": "256GB"})],
      subtitle="Flash drive · USB Type-C", specs=dict(capacity_gb=32, type="Flash drive", interface="USB Type-C"), content="Dual Drive Go"),
    P("sandisk-ultra-dual-drive-luxe-usb-c", "SanDisk", "SanDisk Ultra Dual Drive Luxe USB Type-C",
      [(81, "64GB", {"Capacity": "64GB"}), (82, "128GB", {"Capacity": "128GB"}), (83, "256GB", {"Capacity": "256GB"}), (80, "1TB", {"Capacity": "1TB"})],
      subtitle="Flash drive · USB Type-C + Type-A", specs=dict(capacity_gb=64, type="Flash drive", interface="USB Type-C + USB Type-A"), content="Luxe Dual Drive"),
    P("sandisk-cruzer-blade-cz50", "SanDisk", "SanDisk Cruzer Blade USB Flash Drive (CZ50)",
      [(84, "16GB", {"Capacity": "16GB"}), (85, "128GB", {"Capacity": "128GB"})],
      subtitle="Flash drive · USB", specs=dict(capacity_gb=128, type="Flash drive"), content="CZ50",
      notes="16GB: quantity 0 and no selling price in the stock list, so that variant is switched off until it has a price."),
    P("sandisk-cruzer-glide-cz600", "SanDisk", "SanDisk Cruzer Glide 3.0 USB Flash Drive (CZ600)",
      [(86, "16GB", {"Capacity": "16GB"}), (87, "32GB", {"Capacity": "32GB"}), (88, "128GB", {"Capacity": "128GB"}), (89, "256GB", {"Capacity": "256GB"})],
      subtitle="Flash drive · USB 3.0", specs=dict(capacity_gb=16, type="Flash drive", interface="USB 3.0"), content="CZ600"),
    # ── Networking & components ────────────────────────────────────────────────────────────────
    P("tp-link-tl-wn823n-mini-wireless-usb-adapter", "TP-Link", "TP-Link TL-WN823N Mini Wireless USB Adapter", one(90),
      subtitle="300Mbps · USB Wi-Fi adapter", specs=dict(connectivity=["Wi-Fi 300Mbps"], interface="USB"), content="WN823N"),
    P("kingston-16gb-ddr5-5600-desktop-ram", "Kingston", "Kingston 16GB DDR5 5600 Desktop RAM", one(91),
      subtitle="16GB · DDR5 · 5600 · Desktop", content="Kingston RAM",
      notes="HOLD: no Kingston part number in the stock list — confirm the RAM product line."),
]


# ── workbook ─────────────────────────────────────────────────────────────────────────────────

def money(value):
    """(price, range_text): numbers as-is; '215,000-220,000' → lower figure + the range text."""
    if value is None or value == "":
        return None, None
    if isinstance(value, (int, float)):
        return round(float(value), 2), None
    text = str(value).strip()
    parts = [float(p.replace(",", "")) for p in re.split(r"\s*[-–]\s*", text) if p.strip()]
    return (min(parts), text) if parts else (None, None)


def read_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    inv = {}
    for i, r in enumerate(wb["Combined Inventory"].iter_rows(min_row=6, values_only=True), start=6):
        if not any(v is not None for v in r):
            continue
        category, item, qty, cost, dealer, final, colour, other, source = r[:9]
        price, price_range = money(final)
        inv[i] = dict(category=category, item=item, qty=qty, cost=money(cost)[0], dealer=money(dealer)[0],
                      price=price, price_range=price_range, colour=colour, other=other, source=source)
    content = {}
    for r in wb["Brand Content"].iter_rows(min_row=6, values_only=True):
        if r[0]:
            content[r[0]] = dict(title=r[1], excerpt=r[2], wording=r[3], url=r[4], notes=r[5])
    return inv, content


# ── SQL ──────────────────────────────────────────────────────────────────────────────────────

def q(value):
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return f"{value:.2f}".rstrip("0").rstrip(".") if isinstance(value, float) else str(value)
    return "'" + str(value).replace("'", "''") + "'"


def qa(items):
    return "ARRAY[" + ", ".join(q(i) for i in items) + "]::TEXT[]" if items else "'{}'::TEXT[]"


def qj(obj):
    return q(json.dumps(obj, ensure_ascii=False)) + "::jsonb"


def build(inv, content):
    used = sorted(r for p in PRODUCTS for r, _, _ in p["rows"])
    assert used == sorted(inv), f"rows not covered exactly once: missing {sorted(set(inv) - set(used))}, extra {sorted(set(used) - set(inv))}"
    assert len(set(p["slug"] for p in PRODUCTS)) == len(PRODUCTS), "duplicate slug"

    products, variants, stock, psrc, vsrc, costs = [], [], [], [], [], []
    for order, p in enumerate(PRODUCTS, start=1):
        rows = [inv[r] for r, _, _ in p["rows"]]
        cats = {row["category"] for row in rows}
        depts = {DEPARTMENTS[c] for c in cats}
        assert len(depts) == 1, (p["slug"], depts)
        brand = content.get(p["content"]) if p["content"] else None
        assert p["content"] is None or brand, f"{p['slug']}: no Brand Content row '{p['content']}'"
        attributes = {}
        if p["specs"]:
            attributes["specs"] = p["specs"]
        if brand and brand["wording"] == "Short verbatim excerpt" and brand["excerpt"]:
            excerpt = brand["excerpt"].strip()
            attributes["highlights"] = [excerpt[0].upper() + excerpt[1:]]
        if p["use_cases"]:
            attributes["use_cases"] = p["use_cases"]
        tags = sorted({TAGS.get(c, c.lower()) for c in cats})
        products.append((p["slug"], p["brand"], p["name"], p["subtitle"], depts.pop(), tags, attributes, order * 10))
        if brand or p["notes"]:
            notes = " ".join(x for x in [brand["notes"] if brand else None, p["notes"]] if x)
            psrc.append((p["slug"], brand["title"] if brand else None, brand["url"] if brand else None,
                         brand["wording"] if brand else None, notes or None))
        for position, (r, vname, options) in enumerate(p["rows"]):
            row = inv[r]
            sellable = bool(row["price"]) and row["price"] > 0
            variants.append((p["slug"], vname, options, row["price"] or 0, position, sellable))
            if row["qty"] is not None:
                stock.append((p["slug"], vname, int(row["qty"])))
            note = None
            if row["price_range"]:
                note = f"Final selling range in the stock list: {row['price_range']} (shop price set to the lower figure)."
            elif not sellable:
                note = "No selling price in the stock list — variant switched off until a price is set."
            vsrc.append((p["slug"], vname, r, row["category"], row["source"] or row["item"], row["colour"], row["other"], note))
            costs.append((p["slug"], vname, row["cost"], row["dealer"] or None))  # a 0 dealer price means "not set"
    return products, variants, stock, psrc, vsrc, costs


def values(rows, fmt):
    return ",\n".join("      (" + fmt(r) + ")" for r in rows)


HEADER = """-- ═════════════════════════════════════════════════════════════════════════════
-- {name} — Dock One Solutions
-- GENERATED by scripts/db/inventory/build_seed.py from the client's handoff workbook
-- Dock_One_Combined_Inventory.xlsx. Do not hand-edit — change the script and re-run it.
--
{body}
-- ═════════════════════════════════════════════════════════════════════════════
"""


def seed_sql(products, variants, stock, psrc, vsrc):
    n_products, n_variants = len(products), len(variants)
    n_active = sum(1 for v in variants if v[5])
    n_stock = sum(s[2] for s in stock)
    slugs = ",\n    ".join(", ".join(q(p[0]) for p in products[i:i + 3]) for i in range(0, len(products), 3))
    demo = ",\n    ".join(", ".join(q(s) for s in DEMO_SLUGS[i:i + 4]) for i in range(0, len(DEMO_SLUGS), 4))
    body = f"""-- PURPOSE      The store's REAL catalogue: {n_products} products / {n_variants} variants from the workbook's
--              86 stock rows, with selling prices, stock levels and the brand-match notes.
--                · 5 new storefront departments (the sheet's 20 categories grouped — approved), the
--                  four existing ones kept; "Mice" becomes "Mice & Mousepads" if still unrenamed.
--                · rows that differ only by colour / size / capacity / switch / GPU are ONE product
--                  with priced variants (option_values); each variant has its own stock.
--                · price = "Final Selling"; a range ("215,000-220,000") uses the LOWER figure.
--                · stock = "Quantity" (an inventory row per variant, low-stock warning at 3).
--                · the 16 DEMO products and their 4 demo collections are switched OFF
--                  (is_active = FALSE), not deleted — turn any back on in the admin if wanted.
--              Not filled (left for the owner — nothing invented): photos, descriptions, warranty,
--              SKUs, category stage images/hero products, merchandising flags (new / best seller /
--              flash deal / featured). Product pages show the brand's short verbatim excerpt as a
--              highlight where the workbook has one.
-- ADDS         two ADMIN-ONLY tables for what the workbook has and the schema had no place for:
--                product_sourcing  (per product)  official brand title, brand page URL, wording
--                                                  type, review notes (incl. HOLD items)
--                variant_sourcing  (per variant)  workbook row, sheet category, original source
--                                                  description, colour / other-variant text,
--                                                  dealer price (filled by 35), price note
--              Kept out of product_costs on purpose: the admin's "clear cost" deletes that row.
-- DEPENDS ON   01, 02 (is_admin), 04, 05, 30 (the demo slugs it hides).
-- ENABLES      35_seed_inventory_costs (cost + dealer prices, kept out of git).
-- RULES        rows by slug + variant name only (never SERIAL ids); ON CONFLICT DO NOTHING; one
--              atomic DO block behind the app_config marker 'seed_34_inventory'; self-verifying.
-- SAFE TO RE-RUN: yes (second run is a no-op; the demo products are hidden only on the first run)."""
    sql = [HEADER.format(name=OUT_SEED.name, body=body)]
    sql.append("""
-- ─────────────────────────────────────────────────────────────────────────────
-- 1. Admin-only sourcing tables (the workbook fields the schema had no column for)
-- ─────────────────────────────────────────────────────────────────────────────

CREATE TABLE IF NOT EXISTS public.product_sourcing (
  product_id     INT PRIMARY KEY REFERENCES public.products (id) ON DELETE CASCADE,
  official_title TEXT,                      -- the brand's own product title / model label
  official_url   TEXT,                      -- the brand page the match was made against
  wording_type   TEXT,                      -- workbook: how far the brand wording was verified
  review_notes   TEXT,                      -- engineer notes; "HOLD:" = confirm before publishing
  updated_at     TIMESTAMPTZ NOT NULL DEFAULT now(),
  CONSTRAINT product_sourcing_lengths CHECK (
        (official_title IS NULL OR char_length(official_title) <= 300)
    AND (official_url   IS NULL OR (char_length(official_url) <= 2000 AND official_url ~ '^https://'))
    AND (wording_type   IS NULL OR char_length(wording_type) <= 120)
    AND (review_notes   IS NULL OR char_length(review_notes) <= 4000))
);

CREATE TABLE IF NOT EXISTS public.variant_sourcing (
  variant_id         INT PRIMARY KEY REFERENCES public.product_variants (id) ON DELETE CASCADE,
  source_row         INT,                   -- row number in the workbook's "Combined Inventory" sheet
  source_category    TEXT,                  -- the sheet's category (e.g. "Gaming Combos", "Internal SSD")
  source_description TEXT,                  -- the original stock-list text, verbatim
  colour             TEXT,                  -- the sheet's Color column ("Not specified" kept as-is)
  other_variant      TEXT,                  -- the sheet's Other variant column
  dealer_price       NUMERIC(12,2),         -- "Dealer Selling" (35_seed_inventory_costs)
  price_note         TEXT,                  -- e.g. the original selling-price range
  updated_at         TIMESTAMPTZ NOT NULL DEFAULT now(),
  CONSTRAINT variant_sourcing_dealer_valid CHECK (dealer_price IS NULL OR dealer_price >= 0),
  CONSTRAINT variant_sourcing_lengths CHECK (
        (source_category    IS NULL OR char_length(source_category)    <= 120)
    AND (source_description IS NULL OR char_length(source_description) <= 1000)
    AND (colour             IS NULL OR char_length(colour)             <= 120)
    AND (other_variant      IS NULL OR char_length(other_variant)      <= 200)
    AND (price_note         IS NULL OR char_length(price_note)         <= 1000))
);

DROP TRIGGER IF EXISTS product_sourcing_touch ON public.product_sourcing;
CREATE TRIGGER product_sourcing_touch BEFORE UPDATE ON public.product_sourcing
  FOR EACH ROW EXECUTE FUNCTION public.touch_updated_at();
DROP TRIGGER IF EXISTS variant_sourcing_touch ON public.variant_sourcing;
CREATE TRIGGER variant_sourcing_touch BEFORE UPDATE ON public.variant_sourcing
  FOR EACH ROW EXECUTE FUNCTION public.touch_updated_at();

-- Same seal as product_costs: never anon, admins only through RLS, nobody TRUNCATEs via the API.
ALTER TABLE public.product_sourcing ENABLE ROW LEVEL SECURITY;
ALTER TABLE public.variant_sourcing ENABLE ROW LEVEL SECURITY;
REVOKE ALL ON public.product_sourcing, public.variant_sourcing FROM anon;
REVOKE TRUNCATE, REFERENCES, TRIGGER ON public.product_sourcing, public.variant_sourcing FROM authenticated;
DROP POLICY IF EXISTS product_sourcing_admin_all ON public.product_sourcing;
CREATE POLICY product_sourcing_admin_all ON public.product_sourcing
  FOR ALL TO authenticated USING ((SELECT public.is_admin())) WITH CHECK ((SELECT public.is_admin()));
DROP POLICY IF EXISTS variant_sourcing_admin_all ON public.variant_sourcing;
CREATE POLICY variant_sourcing_admin_all ON public.variant_sourcing
  FOR ALL TO authenticated USING ((SELECT public.is_admin())) WITH CHECK ((SELECT public.is_admin()));
""")
    cats = ",\n".join(f"    ({q(k)}, {q(v[0])}, {q(v[1])}, {q(v[2])}, {v[3]})" for k, v in NEW_CATEGORIES.items())
    sql.append(f"""
-- ─────────────────────────────────────────────────────────────────────────────
-- 2. The catalogue
-- ─────────────────────────────────────────────────────────────────────────────

DO $seed$
DECLARE
  c_marker CONSTANT TEXT := 'seed_34_inventory';
  c_slugs  CONSTANT TEXT[] := ARRAY[
    {slugs}];
  c_demo   CONSTANT TEXT[] := ARRAY[
    {demo}];
  c_demo_collections CONSTANT TEXT[] := ARRAY['work-from-home', 'gaming-zone', 'campus-kit', 'creator-studio'];
  v_bad    TEXT;
  v_n      INT;
  v_sum    BIGINT;
BEGIN
  IF EXISTS (SELECT 1 FROM public.app_config WHERE name = c_marker) THEN
    RAISE NOTICE '34_seed_inventory: already applied (app_config %) — skipped', c_marker;
    RETURN;
  END IF;

  -- ── 2a. Departments: five new ones; the four existing keep the owner's copy ─────────────────
  INSERT INTO public.categories (id, name, tagline, scene, sort_order)
  VALUES
{cats}
  ON CONFLICT (id) DO NOTHING;
  UPDATE public.categories SET name = 'Mice & Mousepads', tagline = 'Gaming & wireless mice, mousepads'
   WHERE id = 'mice' AND name = 'Mice';

  -- ── 2b. The demo catalogue goes dark (reversible: admin → Products → Active) ────────────────
  UPDATE public.products SET is_active = FALSE WHERE slug = ANY (c_demo) AND is_active;
  UPDATE public.collections SET is_active = FALSE WHERE id = ANY (c_demo_collections) AND is_active;

  -- ── 2c. Products (price columns derive from the variants in 2d) ─────────────────────────────
  INSERT INTO public.products (slug, brand, name, subtitle, category_id, tags, attributes, sort_order)
  SELECT d.slug, d.brand, d.name, d.subtitle, d.category_id, d.tags, d.attributes, d.sort_order
    FROM (VALUES
{values(products, lambda p: ", ".join([q(p[0]), q(p[1]), q(p[2]), q(p[3]), q(p[4]), qa(p[5]), qj(p[6]), str(p[7])]))}
    ) AS d(slug, brand, name, subtitle, category_id, tags, attributes, sort_order)
  ON CONFLICT (slug) DO NOTHING;

  -- ── 2d. Variants: one per stock row, priced from "Final Selling" ────────────────────────────
  INSERT INTO public.product_variants (product_id, name, option_values, price, position, is_active)
  SELECT p.id, d.name, d.option_values, d.price, d.position, d.is_active
    FROM (VALUES
{values(variants, lambda v: ", ".join([q(v[0]), q(v[1]), qj(v[2]), q(float(v[3])), str(v[4]), q(v[5])]))}
    ) AS d(slug, name, option_values, price, position, is_active)
    JOIN public.products p ON p.slug = d.slug
  ON CONFLICT DO NOTHING;

  -- ── 2e. Stock: the "Quantity" column, one tracked inventory row per variant ─────────────────
  INSERT INTO public.inventory (variant_id, product_id, stock_level)
  SELECT v.id, v.product_id, d.qty
    FROM (VALUES
{values(stock, lambda s: ", ".join([q(s[0]), q(s[1]), str(s[2])]))}
    ) AS d(slug, name, qty)
    JOIN public.products p ON p.slug = d.slug
    JOIN public.product_variants v ON v.product_id = p.id AND lower(v.name) = lower(d.name)
  ON CONFLICT (variant_id) DO NOTHING;

  -- ── 2f. Sourcing notes (admin-only) ─────────────────────────────────────────────────────────
  INSERT INTO public.product_sourcing (product_id, official_title, official_url, wording_type, review_notes)
  SELECT p.id, d.official_title, d.official_url, d.wording_type, d.review_notes
    FROM (VALUES
{values(psrc, lambda s: ", ".join(q(x) for x in s))}
    ) AS d(slug, official_title, official_url, wording_type, review_notes)
    JOIN public.products p ON p.slug = d.slug
  ON CONFLICT (product_id) DO NOTHING;

  INSERT INTO public.variant_sourcing (variant_id, source_row, source_category, source_description, colour, other_variant, price_note)
  SELECT v.id, d.source_row, d.source_category, d.source_description, d.colour, d.other_variant, d.price_note
    FROM (VALUES
{values(vsrc, lambda s: ", ".join([q(s[0]), q(s[1]), str(s[2])] + [q(x) for x in s[3:]]))}
    ) AS d(slug, name, source_row, source_category, source_description, colour, other_variant, price_note)
    JOIN public.products p ON p.slug = d.slug
    JOIN public.product_variants v ON v.product_id = p.id AND lower(v.name) = lower(d.name)
  ON CONFLICT (variant_id) DO NOTHING;

  INSERT INTO public.app_config (name, value) VALUES (c_marker, to_char(now(), 'YYYY-MM-DD"T"HH24:MI:SSOF'))
  ON CONFLICT (name) DO NOTHING;

  -- ── 2g. Self-verification: any failure rolls the whole seed back ────────────────────────────
  SELECT count(*) INTO v_n FROM public.products WHERE slug = ANY (c_slugs);
  IF v_n <> {n_products} THEN RAISE EXCEPTION '34_seed_inventory: expected {n_products} products, found %', v_n; END IF;

  SELECT count(*) INTO v_n
    FROM public.product_variants v JOIN public.products p ON p.id = v.product_id
   WHERE p.slug = ANY (c_slugs);
  IF v_n <> {n_variants} THEN RAISE EXCEPTION '34_seed_inventory: expected {n_variants} variants, found %', v_n; END IF;

  SELECT count(*) INTO v_n
    FROM public.product_variants v JOIN public.products p ON p.id = v.product_id
   WHERE p.slug = ANY (c_slugs) AND v.is_active AND v.price > 0;
  IF v_n <> {n_active} THEN RAISE EXCEPTION '34_seed_inventory: expected {n_active} priced active variants, found %', v_n; END IF;

  SELECT string_agg(p.slug, ', ') INTO v_bad
    FROM public.products p
   WHERE p.slug = ANY (c_slugs)
     AND (NOT p.is_active OR p.variant_count < 1 OR p.default_variant_id IS NULL OR p.price <= 0 OR p.category_id IS NULL);
  IF v_bad IS NOT NULL THEN RAISE EXCEPTION '34_seed_inventory: products not visible/purchasable: %', v_bad; END IF;

  SELECT count(*), COALESCE(sum(i.stock_level), 0) INTO v_n, v_sum
    FROM public.inventory i JOIN public.products p ON p.id = i.product_id
   WHERE p.slug = ANY (c_slugs);
  IF v_n <> {len(stock)} OR v_sum <> {n_stock} THEN
    RAISE EXCEPTION '34_seed_inventory: expected {len(stock)} stock rows totalling {n_stock}, found % totalling %', v_n, v_sum;
  END IF;

  SELECT count(*) INTO v_n
    FROM public.variant_sourcing s JOIN public.product_variants v ON v.id = s.variant_id
    JOIN public.products p ON p.id = v.product_id
   WHERE p.slug = ANY (c_slugs);
  IF v_n <> {n_variants} THEN RAISE EXCEPTION '34_seed_inventory: expected {n_variants} variant_sourcing rows, found %', v_n; END IF;

  SELECT count(*) INTO v_n FROM public.products WHERE slug = ANY (c_demo) AND is_active;
  IF v_n <> 0 THEN RAISE EXCEPTION '34_seed_inventory: % demo products still active', v_n; END IF;
  SELECT count(*) INTO v_n FROM public.collections WHERE id = ANY (c_demo_collections) AND is_active;
  IF v_n <> 0 THEN RAISE EXCEPTION '34_seed_inventory: % demo collections still active', v_n; END IF;

  RAISE NOTICE '34_seed_inventory: {n_products} products, {n_variants} variants, {len(stock)} stock rows ({n_stock} units); demo catalogue + collections hidden';
END $seed$;

-- ═════════════════════════════════════════════════════════════════════════════
-- OPS NOTE — 34_seed_inventory
--   Next: run 35_seed_inventory_costs.sql (cost + dealer prices; kept out of git), then add photos
--   and descriptions in admin → Products. Review the HOLD notes first:
--     SELECT p.name, s.review_notes FROM public.product_sourcing s JOIN public.products p ON p.id = s.product_id
--      WHERE s.review_notes ILIKE '%HOLD%' OR s.review_notes ILIKE '%confirm%';
--   Price ranges kept from the stock list:
--     SELECT p.name, v.name, v.price, s.price_note FROM public.variant_sourcing s
--       JOIN public.product_variants v ON v.id = s.variant_id JOIN public.products p ON p.id = v.product_id
--      WHERE s.price_note IS NOT NULL ORDER BY p.sort_order, v.position;
--   Bring a demo product / collection back: admin → Products / Collections → Active, or
--     UPDATE public.products SET is_active = TRUE WHERE slug = '<slug>';
-- ═════════════════════════════════════════════════════════════════════════════
""")
    return "".join(sql)


def costs_sql(costs):
    body = """--   ██ CONTAINS THE CLIENT'S COST AND DEALER PRICES — git-ignored; never commit it ██
--
-- PURPOSE      "Cost per unit" → product_costs.cost_price (the admin's margin figures) and
--              "Dealer Selling" → variant_sourcing.dealer_price, for the variants 34 created.
--              A blank cell writes nothing (blank is not zero).
-- DEPENDS ON   34_seed_inventory.
-- RULES        variants by product slug + variant name; marker 'seed_35_inventory_costs';
--              existing cost rows are left alone (the owner's edits win).
-- SAFE TO RE-RUN: yes."""
    rows = [c for c in costs if c[2] is not None or c[3] is not None]
    n_cost = sum(1 for c in costs if c[2] is not None)
    n_dealer = sum(1 for c in costs if c[3] is not None)
    return HEADER.format(name=OUT_COSTS.name, body=body) + f"""
DO $seed$
DECLARE
  c_marker CONSTANT TEXT := 'seed_35_inventory_costs';
  v_n      INT;
BEGIN
  IF EXISTS (SELECT 1 FROM public.app_config WHERE name = c_marker) THEN
    RAISE NOTICE '35_seed_inventory_costs: already applied (app_config %) — skipped', c_marker;
    RETURN;
  END IF;
  IF NOT EXISTS (SELECT 1 FROM public.app_config WHERE name = 'seed_34_inventory') THEN
    RAISE EXCEPTION '35_seed_inventory_costs: run 34_seed_inventory.sql first';
  END IF;

  CREATE TEMP TABLE _costs (slug TEXT, name TEXT, cost NUMERIC(12,2), dealer NUMERIC(12,2)) ON COMMIT DROP;
  INSERT INTO _costs (slug, name, cost, dealer) VALUES
{values(rows, lambda c: ", ".join([q(c[0]), q(c[1]), q(c[2]), q(c[3])]))};

  INSERT INTO public.product_costs (variant_id, cost_price)
  SELECT v.id, c.cost
    FROM _costs c
    JOIN public.products p ON p.slug = c.slug
    JOIN public.product_variants v ON v.product_id = p.id AND lower(v.name) = lower(c.name)
   WHERE c.cost IS NOT NULL
  ON CONFLICT (variant_id) DO NOTHING;

  UPDATE public.variant_sourcing s
     SET dealer_price = c.dealer
    FROM _costs c
    JOIN public.products p ON p.slug = c.slug
    JOIN public.product_variants v ON v.product_id = p.id AND lower(v.name) = lower(c.name)
   WHERE s.variant_id = v.id AND c.dealer IS NOT NULL AND s.dealer_price IS NULL;

  INSERT INTO public.app_config (name, value) VALUES (c_marker, to_char(now(), 'YYYY-MM-DD"T"HH24:MI:SSOF'))
  ON CONFLICT (name) DO NOTHING;

  SELECT count(*) INTO v_n FROM _costs c
    JOIN public.products p ON p.slug = c.slug
    JOIN public.product_variants v ON v.product_id = p.id AND lower(v.name) = lower(c.name)
    JOIN public.product_costs pc ON pc.variant_id = v.id
   WHERE c.cost IS NOT NULL;
  IF v_n <> {n_cost} THEN RAISE EXCEPTION '35_seed_inventory_costs: expected {n_cost} cost rows, found %', v_n; END IF;
  SELECT count(*) INTO v_n FROM _costs c
    JOIN public.products p ON p.slug = c.slug
    JOIN public.product_variants v ON v.product_id = p.id AND lower(v.name) = lower(c.name)
    JOIN public.variant_sourcing s ON s.variant_id = v.id AND s.dealer_price IS NOT NULL
   WHERE c.dealer IS NOT NULL;
  IF v_n <> {n_dealer} THEN RAISE EXCEPTION '35_seed_inventory_costs: expected {n_dealer} dealer prices, found %', v_n; END IF;

  RAISE NOTICE '35_seed_inventory_costs: {n_cost} cost prices, {n_dealer} dealer prices';
END $seed$;
"""


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    inv, content = read_workbook(sys.argv[1])
    products, variants, stock, psrc, vsrc, costs = build(inv, content)
    OUT_SEED.write_text(seed_sql(products, variants, stock, psrc, vsrc))
    OUT_COSTS.write_text(costs_sql(costs))
    print(f"{OUT_SEED.relative_to(ROOT)}: {len(products)} products, {len(variants)} variants, "
          f"{len(stock)} stock rows ({sum(s[2] for s in stock)} units)")
    print(f"{OUT_COSTS.relative_to(ROOT)}: {sum(1 for c in costs if c[2] is not None)} costs, "
          f"{sum(1 for c in costs if c[3] is not None)} dealer prices")


if __name__ == "__main__":
    main()
